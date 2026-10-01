import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_excel():
    wb = openpyxl.Workbook()
    
    # -------------------------------------------------------------
    # Sheet 1: Summary Table (Головна таблиця верифікації)
    # -------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Таблиця верифікації (Task 4)"
    ws1.views.sheetView[0].showGridLines = True

    # Title
    ws1.merge_cells("A1:K1")
    title_cell = ws1["A1"]
    title_cell.value = "РЕЗУЛЬТАТИ ВЕРИФІКАЦІЇ ТРИФАЗНОЇ СМО (ЗАВДАННЯ 4)"
    title_cell.font = Font(name="Calibri", size=15, bold=True, color="1A365D")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 35

    ws1.merge_cells("A2:K2")
    sub_cell = ws1["A2"]
    sub_cell.value = "Схема моделі: CREATE -> PROCESS 1 -> PROCESS 2 -> PROCESS 3 -> DISPOSE | Час моделювання T = 1000"
    sub_cell.font = Font(name="Calibri", size=11, italic=True, color="4A5568")
    sub_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[2].height = 20

    # Headers
    headers = [
        "№",
        "Сценарій експерименту",
        "Інтервал надходження (Create)",
        "Затримка P1",
        "Затримка P2",
        "Затримка P3",
        "Місткість черг (MaxQ)",
        "Створено вимог",
        "Вийшло (Dispose)",
        "Загальна к-ть відмов",
        "Середнє завантаження (P1 / P2 / P3)",
        "Середня довжина черг (P1 / P2 / P3)"
    ]

    header_fill = PatternFill(start_color="2B6CB0", end_color="2B6CB0", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E0'),
        right=Side(style='thin', color='CBD5E0'),
        top=Side(style='thin', color='CBD5E0'),
        bottom=Side(style='thin', color='CBD5E0')
    )

    ws1.row_dimensions[4].height = 28
    for col_num, header in enumerate(headers, 1):
        cell = ws1.cell(row=4, column=col_num)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # Data
    data = [
        [1, "Експ. 1: Базовий стан", 2.0, 1.0, 1.0, 1.0, 5, 500, 482, 16, "0.50 / 0.48 / 0.47", "0.59 / 0.42 / 0.38"],
        [2, "Експ. 2: Високе навантаження", 1.1, 1.0, 1.0, 1.0, 5, 872, 713, 153, "0.78 / 0.72 / 0.76", "1.66 / 1.17 / 1.27"],
        [3, "Експ. 3: 'Вузьке місце' на P2", 2.0, 1.0, 2.2, 1.0, 5, 517, 398, 112, "0.53 / 0.89 / 0.39", "0.44 / 2.65 / 0.25"],
        [4, "Експ. 4: Мала черга (MaxQ=1)", 1.5, 1.3, 1.3, 1.3, 1, 693, 357, 333, "0.63 / 0.54 / 0.51", "0.29 / 0.23 / 0.19"],
        [5, "Експ. 5: Низьке навантаження", 3.0, 0.8, 0.8, 0.8, 5, 335, 334, 0, "0.25 / 0.26 / 0.27", "0.09 / 0.11 / 0.13"]
    ]

    zebra_fill = PatternFill(start_color="F7FAFC", end_color="F7FAFC", fill_type="solid")
    regular_font = Font(name="Calibri", size=11, color="2D3748")
    bold_font = Font(name="Calibri", size=11, bold=True, color="2D3748")

    for row_idx, row_data in enumerate(data, 5):
        ws1.row_dimensions[row_idx].height = 24
        is_zebra = (row_idx % 2 == 1)
        for col_idx, value in enumerate(row_data, 1):
            cell = ws1.cell(row=row_idx, column=col_idx)
            cell.value = value
            cell.font = bold_font if col_idx in [1, 2, 8, 9, 10] else regular_font
            cell.border = thin_border
            if is_zebra:
                cell.fill = zebra_fill
            
            # Alignments
            if col_idx == 2:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    # Explanations below table
    note_row = 12
    ws1.cell(row=note_row, column=2).value = "Висновки та аналіз результатів верифікації:"
    ws1.cell(row=note_row, column=2).font = Font(name="Calibri", size=12, bold=True, color="1A365D")
    
    notes = [
        "1. Експ. 1 (Базовий): При інтервалі надходження 2.0 та обслуговуванні 1.0 на кожній фазі навантаження становить ~50%, втрати мінімальні (16 відмов або 3.2%).",
        "2. Експ. 2 (Високе навантаження): Зменшення інтервалу до 1.1 призводить до різкого стрибка завантаження пристроїв до 72-78% та зростання середньої черги до 1.66.",
        "3. Експ. 3 (Вузьке місце P2): Збільшення затримки на P2 до 2.2 перевантажує даний прилад до 89.1%, черга перед ним зростає до 2.65, і майже всі відмови концентруються на ньому.",
        "4. Експ. 4 (Обмежена черга): При MaxQ=1 система втрачає здатність згладжувати випадкові згущення потоку, що викликає сплеск втрат заявок (333 відмови із 693 створених).",
        "5. Експ. 5 (Низьке навантаження): При низькому вхідному потоці (інтервал 3.0) завантаження пристроїв падає до 25%, відмови повністю відсутні (0 відмов)."
    ]
    for idx, note in enumerate(notes, note_row + 1):
        c = ws1.cell(row=idx, column=2)
        c.value = note
        c.font = Font(name="Calibri", size=10.5, color="4A5568")

    # -------------------------------------------------------------
    # Sheet 2: Детальні результати по кожному блоку (P1, P2, P3)
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="Деталізація по фазах")
    ws2.views.sheetView[0].showGridLines = True

    ws2.merge_cells("A1:K1")
    t2 = ws2["A1"]
    t2.value = "ДЕТАЛЬНІ ПОКАЗНИКИ ОБСЛУГОВУВАННЯ ПО КОЖНІЙ ФАЗІ СМО"
    t2.font = Font(name="Calibri", size=14, bold=True, color="1A365D")
    t2.alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 30

    h2 = [
        "Сценарій", "Блок", "Затримка", "Оброблено вимог", "Середня черга",
        "К-ть відмов", "Ймовірність відмови", "Завантаження каналу"
    ]
    ws2.row_dimensions[3].height = 25
    for c_idx, h in enumerate(h2, 1):
        cell = ws2.cell(row=3, column=c_idx)
        cell.value = h
        cell.font = header_font
        cell.fill = PatternFill(start_color="3182CE", end_color="3182CE", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    detailed_data = [
        ["Експ. 1: Базовий", "PROCESSOR 1", 1.0, 492, 0.5885, 7, 0.0140, 0.5014],
        ["Експ. 1: Базовий", "PROCESSOR 2", 1.0, 490, 0.4230, 2, 0.0041, 0.4750],
        ["Експ. 1: Базовий", "PROCESSOR 3", 1.0, 482, 0.3823, 7, 0.0143, 0.4707],

        ["Експ. 2: Високе навантаження", "PROCESSOR 1", 1.0, 784, 1.6647, 88, 0.1009, 0.7819],
        ["Експ. 2: Високе навантаження", "PROCESSOR 2", 1.0, 741, 1.1718, 38, 0.0488, 0.7202],
        ["Експ. 2: Високе навантаження", "PROCESSOR 3", 1.0, 713, 1.2671, 27, 0.0365, 0.7610],

        ["Експ. 3: Вузьке місце P2", "PROCESSOR 1", 1.0, 513, 0.4341, 3, 0.0058, 0.5263],
        ["Експ. 3: Вузьке місце P2", "PROCESSOR 2", 2.2, 402, 2.6405, 105, 0.2071, 0.8909],
        ["Експ. 3: Вузьке місце P2", "PROCESSOR 3", 1.0, 398, 0.2491, 4, 0.0100, 0.3887],

        ["Експ. 4: Мала черга (Q=1)", "PROCESSOR 1", 1.3, 494, 0.2924, 198, 0.2861, 0.6299],
        ["Експ. 4: Мала черга (Q=1)", "PROCESSOR 2", 1.3, 413, 0.2265, 80, 0.1623, 0.5374],
        ["Експ. 4: Мала черга (Q=1)", "PROCESSOR 3", 1.3, 357, 0.1919, 55, 0.1335, 0.5074],

        ["Експ. 5: Низьке навантаження", "PROCESSOR 1", 0.8, 334, 0.0941, 0, 0.0000, 0.2450],
        ["Експ. 5: Низьке навантаження", "PROCESSOR 2", 0.8, 334, 0.1128, 0, 0.0000, 0.2590],
        ["Експ. 5: Низьке навантаження", "PROCESSOR 3", 0.8, 334, 0.1295, 0, 0.0000, 0.2731]
    ]

    for r_idx, row_vals in enumerate(detailed_data, 4):
        ws2.row_dimensions[r_idx].height = 20
        is_zebra = ((r_idx // 3) % 2 == 1)
        for c_idx, val in enumerate(row_vals, 1):
            cell = ws2.cell(row=r_idx, column=c_idx)
            cell.value = val
            cell.font = regular_font
            cell.border = thin_border
            if is_zebra:
                cell.fill = zebra_fill

            if c_idx in [1, 2]:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif c_idx in [7, 8]:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "0.00%"
            elif c_idx in [5]:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "0.0000"
            else:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    # Auto-adjust column widths
    for sheet in [ws1, ws2]:
        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                # Skip merged title rows
                if cell.row in [1, 2, 12, 13, 14, 15, 16, 17]:
                    continue
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

    ws1.column_dimensions['A'].width = 6
    ws1.column_dimensions['B'].width = 30
    ws1.column_dimensions['K'].width = 34
    ws1.column_dimensions['L'].width = 34

    file_name = "Task4_Verification_Results.xlsx"
    wb.save(file_name)
    print(f"Excel file created: {file_name}")

if __name__ == "__main__":
    generate_excel()
