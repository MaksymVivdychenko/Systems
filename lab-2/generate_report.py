import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_report():
    doc = Document()

    # Page Margins (2 cm)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    def set_cell_shading(cell, color_hex):
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    def add_title(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
        p.paragraph_format.space_after = Pt(6)

    def add_subtitle(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
        p.paragraph_format.space_after = Pt(16)

    def add_h1(text):
        h = doc.add_heading(level=1)
        run = h.add_run(text)
        run.bold = True
        run.font.name = 'Times New Roman'
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)

    def add_h2(text):
        h = doc.add_heading(level=2)
        run = h.add_run(text)
        run.bold = True
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12.5)
        run.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(3)

    def add_code_block(code_str):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.8)
        set_cell_shading(cell, "F7FAFC")
        set_cell_margins(cell, 120, 120, 180, 180)
        
        # Border
        borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E0"/><w:left w:val="single" w:sz="18" w:space="0" w:color="3182CE"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E0"/><w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E0"/></w:tcBorders>')
        cell._tc.get_or_add_tcPr().append(borders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(code_str.strip())
        run.font.name = 'Consolas'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # --- Title & Header ---
    add_title("КОМП’ЮТЕРНИЙ ПРАКТИКУМ № 2")
    add_subtitle("Тема: Об’єктно-орієнтований підхід до побудови імітаційних моделей дискретно-подійних систем\nДисципліна: Моделювання систем")

    # --- Мета роботи ---
    add_h1("1. Мета роботи")
    p = doc.add_paragraph()
    p.add_run("Ознайомлення з принципами об’єктно-орієнтованого підходу до побудови імітаційних моделей дискретно-подійних систем (ДПС). Опанування механізму просування модельного часу за принципом найближчої події. Програмна реалізація, розширення та верифікація імітаційної моделі системи масового обслуговування мовою C# з підтримкою багатоканальності та складної топології маршрутизації (включаючи зворотні зв'язки).")

    # --- Завдання ---
    add_h1("2. Завдання до роботи")
    tasks = [
        "1. Реалізувати алгоритм імітації простої моделі обслуговування одним пристроєм з використанням об’єктно-орієнтованого підходу (мова C#).",
        "2. Модифікувати алгоритм, додавши обчислення середнього завантаження пристрою.",
        "3. Створити модель послідовної СМО за схемою: CREATE -> PROCESS 1 -> PROCESS 2 -> PROCESS 3 -> DISPOSE.",
        "4. Виконати верифікацію моделі за завданням 3, змінюючи значення вхідних змінних та параметрів моделі. Навести результати верифікації у таблиці.",
        "5. Модифікувати клас PROCESS, щоб можна було його використовувати для моделювання процесу обслуговування кількома ідентичними пристроями (багатоканальний пристрій).",
        "6. Модифікувати клас PROCESS, щоб можна було організовувати вихід у два і більше наступних блоків, зокрема з поверненням у попередні блоки (розгалуження та зворотний зв’язок)."
    ]
    for t in tasks:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(t)

    # --- Архітектура класів ---
    add_h1("3. Об’єктно-орієнтована архітектура моделі")
    p = doc.add_paragraph()
    p.add_run("Програма побудована на базі компонентно-орієнтованої архітектури дискретно-подійного моделювання (DES) з модульним розподілом класів по окремих файлах:")
    
    classes_desc = [
        ("Element (Element.cs)", "Базовий абстрактний елемент системи. Містить поля стану (state), модельного часу (tcurr), часу наступної події (tnext), середньої затримки (delayMean), типу розподілу (exp, norm, unif) та віртуальні методи життєвого циклу InAct(), OutAct(), DoStatistics(delta)."),
        ("Create (Create.cs)", "Джерело вимог. Генерує надходження нових заявок за заданим законом розподілу та передає їх на наступний елемент через InAct()."),
        ("Process (Process.cs)", "Обслуговуючий пристрій (СМО). Реалізує накопичувач (чергу) обмеженої або необмеженої ємності, багатоканальне обслуговування (N пристроїв), а також маршрутизацію вихідних потоків за ймовірностями."),
        ("Dispose (Dispose.cs)", "Кінцевий стік системи (Sink). Фіксує кількість заявок, що успішно пройшли всі етапи обслуговування та покинули систему."),
        ("Route (Route.cs)", "Клас опису маршруту. Інкапсулює посилання на цільовий Element, ймовірність переходу (Probability) та лічильник фактично здійснених переходів (TransitionCount)."),
        ("Model (Model.cs)", "Диспетчер модельного часу. Реалізує цикл імітації за принципом найближчої події (Next-Event Time Advance), просування глобального часу tcurr = min(tnext), збір часових інтегральних статистик та виведення фінальних звітів."),
        ("FunRand (FunRand.cs)", "Генератор псевдовипадкових величин за експоненційним, рівномірним та нормальним (метод Бокса-Мюллера) законами розподілу.")
    ]
    for name, desc in classes_desc:
        p = doc.add_paragraph(style='List Bullet')
        r = p.add_run(name + ": ")
        r.bold = True
        p.add_run(desc)

    # --- Механізм просування часу ---
    add_h1("4. Принцип функціонування та просування модельного часу")
    p = doc.add_paragraph()
    p.add_run("У даному практикумі реалізовано ")
    r = p.add_run("принцип найближчої події (Next-Event Time Advance)")
    r.bold = True
    p.add_run(". На відміну від крокового принципу (Δt = const), модельний час не протікає безперервно, а здійснює миттєві стрибки (телепортацію) від поточного моменту tcurr безпосередньо до найближчої майбутньої події:")
    
    add_code_block("""// Фрагмент Model.Simulate:
tnext = Double.MaxValue;
foreach (Element e in list) {
    if (e.GetTnext() < tnext) {
        tnext = e.GetTnext(); // Пошук найближчої події
        eventId = e.GetId();
    }
}
// Інтегрування статистик на інтервалі дельта t:
foreach (Element e in list) e.DoStatistics(tnext - tcurr);

tcurr = tnext; // Миттєве просування модельного часу
list[eventIndex].OutAct(); // Виконання події, що настала""")

    # --- Результати за завданнями 1-2 ---
    add_h1("5. Реалізація завдань 1 та 2 (Одноканальна СМО та розрахунок завантаження)")
    p = doc.add_paragraph()
    p.add_run("Реалізовано базову модель обслуговування CREATE -> PROCESS з ємністю черги maxqueue = 5. Додано обчислення середнього завантаження пристрою (коефіцієнта використання) шляхом інтегрування стану зайнятості пристрою за часом:")
    
    p_eq = doc.add_paragraph()
    p_eq.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_eq = p_eq.add_run("K_завантаження = ( ∫ state(t) dt ) / T_моделювання = meanLoad / tcurr")
    r_eq.bold = True

    add_code_block("""// Збір статистики у Process.DoStatistics:
meanQueue = GetMeanQueue() + queue * delta;
meanLoad = meanLoad + GetBusyChannelsCount() * delta;""")

    p = doc.add_paragraph()
    p.add_run("Результати моделювання на інтервалі T = 1000:")
    
    # Task 1-2 table
    tbl1 = doc.add_table(rows=6, cols=2)
    tbl1.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = [("Параметр", "Значення"),
               ("Кількість створених вимог (CREATOR)", "537"),
               ("Кількість обслугованих вимог (PROCESSOR)", "530"),
               ("Середня довжина черги", "0.5932"),
               ("Ймовірність відмови (втрати вимоги)", "0.0094 (0.94%)"),
               ("Середнє завантаження пристрою", "0.5423 (54.23%)")]
    for i, (k, v) in enumerate(headers):
        cell_k, cell_v = tbl1.cell(i, 0), tbl1.cell(i, 1)
        cell_k.width, cell_v.width = Inches(3.4), Inches(3.4)
        set_cell_margins(cell_k, 60, 60, 100, 100)
        set_cell_margins(cell_v, 60, 60, 100, 100)
        cell_k.paragraphs[0].add_run(k).bold = (i == 0)
        cell_v.paragraphs[0].add_run(v).bold = (i == 0)
        if i == 0:
            set_cell_shading(cell_k, "E2E8F0")
            set_cell_shading(cell_v, "E2E8F0")

    # --- Завдання 3 ---
    add_h1("6. Реалізація завдання 3 (Трифазна послідовна СМО)")
    p = doc.add_paragraph()
    p.add_run("Створено імітаційну модель послідовного ланцюга відповідно до схеми рисунка 2.1 методичних вказівок:")
    p_sc = doc.add_paragraph()
    p_sc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sc = p_sc.add_run("[ CREATE ]  --->  [ PROCESS 1 ]  --->  [ PROCESS 2 ]  --->  [ PROCESS 3 ]  --->  [ DISPOSE ]")
    r_sc.bold = True
    r_sc.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)

    p = doc.add_paragraph()
    p.add_run("Для передачі заявок між послідовними блоками метод Process.OutAct() викликає base.GetNextElement().InAct(). Результати роботи моделі при delayCreate = 2.0, delayP1 = 1.0, delayP2 = 1.0, delayP3 = 1.0, maxqueue = 5:")

    tbl3 = doc.add_table(rows=6, cols=5)
    tbl3.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers3 = ["Показник", "CREATOR", "PROCESSOR 1", "PROCESSOR 2", "PROCESSOR 3"]
    for j, h in enumerate(headers3):
        c = tbl3.cell(0, j)
        set_cell_shading(c, "E2E8F0")
        set_cell_margins(c, 60, 60, 80, 80)
        c.paragraphs[0].add_run(h).bold = True

    data3 = [
        ["Оброблено вимог", "517", "513", "512", "505 (у DISPOSE)"],
        ["Середня черга", "-", "0.4469", "0.3702", "0.5458"],
        ["Кількість відмов", "-", "4", "1", "4"],
        ["Ймовірність відмови", "-", "0.0078", "0.0020", "0.0079"],
        ["Коефіцієнт завантаження", "-", "0.4899", "0.4962", "0.5441"]
    ]
    for i, row in enumerate(data3):
        for j, val in enumerate(row):
            c = tbl3.cell(i+1, j)
            set_cell_margins(c, 50, 50, 80, 80)
            c.paragraphs[0].add_run(val)

    # --- Завдання 4: Верифікація ---
    add_h1("7. Реалізація завдання 4 (Верифікація моделі з варіюванням параметрів)")
    p = doc.add_paragraph()
    p.add_run("Проведено серію із 5 експериментів на модельний час T = 1000 для перевірки адекватності реакції системи на зміну інтенсивності вхідного потоку, тривалості обслуговування на фазах та ємності накопичувачів:")

    tbl4 = doc.add_table(rows=6, cols=7)
    tbl4.alignment = WD_TABLE_ALIGNMENT.CENTER
    h4 = ["№ / Сценарій", "Затримки (C / P1, P2, P3)", "MaxQ", "Створено / Вийшло", "Відмови (P1+P2+P3)", "Завантаження (P1 / P2 / P3)", "Черги (P1 / P2 / P3)"]
    for j, h in enumerate(h4):
        c = tbl4.cell(0, j)
        set_cell_shading(c, "E2E8F0")
        set_cell_margins(c, 60, 60, 60, 60)
        run = c.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(10)

    data4 = [
        ["1. Базовий стан", "2.0 / (1.0, 1.0, 1.0)", "5", "504 / 485", "16", "0.52 / 0.48 / 0.47", "0.49 / 0.41 / 0.41"],
        ["2. Високе навантаження", "1.1 / (1.0, 1.0, 1.0)", "5", "877 / 722", "150", "0.82 / 0.73 / 0.74", "1.49 / 1.32 / 1.38"],
        ["3. 'Вузьке місце' (P2)", "2.0 / (1.0, 2.2, 1.0)", "5", "502 / 405", "91 (P2=88)", "0.49 / 0.94 / 0.42", "0.36 / 2.66 / 0.31"],
        ["4. Мала черга (Q=1)", "1.5 / (1.3, 1.3, 1.3)", "1", "673 / 362", "307", "0.60 / 0.54 / 0.48", "0.29 / 0.19 / 0.17"],
        ["5. Низьке навантаження", "3.0 / (0.8, 0.8, 0.8)", "5", "307 / 306", "0", "0.24 / 0.24 / 0.25", "0.06 / 0.06 / 0.07"]
    ]
    for i, row in enumerate(data4):
        for j, val in enumerate(row):
            c = tbl4.cell(i+1, j)
            set_cell_margins(c, 50, 50, 60, 60)
            run = c.paragraphs[0].add_run(val)
            run.font.size = Pt(9.5)

    p_v = doc.add_paragraph()
    p_v.paragraph_format.space_before = Pt(6)
    p_v.add_run("Висновки з верифікації:\n").bold = True
    v_notes = [
        "Збільшення інтенсивності вхідного потоку (Експ. 2) призводить до закономірного зростання завантаження пристроїв до 73-82% та збільшення відмов через переповнення черг.",
        "Штучне створення 'вузького місця' на фазі P2 (Експ. 3) завантажує даний пристрій до 94.2%, черга перед ним зростає вшестеро (до 2.66), і практично всі відмови системи (88 із 91) концентруються на ньому.",
        "Зменшення місткості черги до 1 вимоги (Експ. 4) викликає сплеск відмов (307 вимог втрачено), оскільки система позбавлена можливості згладжувати випадкові сплески надходжень.",
        "У ненапруженому режимі (Експ. 5) відмови повністю відсутні, завантаження становить ~24%, що підтверджує коректність роботи алгоритму."
    ]
    for vn in v_notes:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(vn)

    # --- Завдання 5: Багатоканальність ---
    add_h1("8. Реалізація завдання 5 (Багатоканальний обслуговуючий пристрій)")
    p = doc.add_paragraph()
    p.add_run("Клас Process модифіковано для підтримки довільної кількості паралельних ідентичних каналів обслуговування (приладів). Стан кожного каналу та індивідуальні моменти звільнення відстежуються у масивах channelStates[channels] та tnextChannels[channels]. Базове поле tnext елемента завжди відповідає найближчому завершенню серед зайнятих каналів.")
    
    add_h2("Теоретична верифікація за формулами теорії масового обслуговування (СМО M/M/c):")
    p = doc.add_paragraph()
    p.add_run("Для моделі з інтенсивністю надходження λ = 1.0 (mean = 1.0), тривалістю обслуговування на канал 1/μ = 2.0 (μ = 0.5) та c = 3 каналами теоретичний коефіцієнт завантаження каналу дорівнює:")
    
    p_f = doc.add_paragraph()
    p_f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_f = p_f.add_run("ρ_теор = λ / (c · μ) = 1.0 / (3 · 0.5) = 2/3 ≈ 0.6667 (66.67%)")
    r_f.bold = True

    p = doc.add_paragraph()
    p.add_run("Порівняння теоретичних та імітаційних результатів:")
    
    tbl5 = doc.add_table(rows=5, cols=3)
    tbl5.alignment = WD_TABLE_ALIGNMENT.CENTER
    h5 = ["Параметр", "Теоретичне значення", "Отримане в моделі"]
    for j, h in enumerate(h5):
        c = tbl5.cell(0, j)
        set_cell_shading(c, "E2E8F0")
        set_cell_margins(c, 60, 60, 80, 80)
        c.paragraphs[0].add_run(h).bold = True
    
    d5 = [
        ["Кількість каналів (c)", "3", "3"],
        ["Коефіцієнт завантаження каналу (ρ)", "0.6667 (66.7%)", "0.6838 (68.4%)"],
        ["Середня кількість зайнятих каналів", "2.0000", "2.0513 з 3"],
        ["Баланс вимог: Створено = Оброблено + Відмови + В системі", "1060", "1028 + 27 + 5 = 1060 (Точний збіг)"]
    ]
    for i, row in enumerate(d5):
        for j, val in enumerate(row):
            c = tbl5.cell(i+1, j)
            set_cell_margins(c, 50, 50, 80, 80)
            c.paragraphs[0].add_run(val)

    # --- Завдання 6: Маршрутизація та зворотний зв'язок ---
    add_h1("9. Реалізація завдання 6 (Розгалуження та зворотний зв'язок)")
    p = doc.add_paragraph()
    p.add_run("Створено структуру маршрутизації Route(NextElement, Probability, TransitionCount). Вибір напрямку виходу реалізовано за методом кумулятивних ймовірностей (рулетки). Забезпечено можливість адресації попередніх блоків ланцюга для організації циклів повторного обслуговування (зворотного зв'язку).")
    
    add_h2("Топологія верифікаційної моделі завдання 6:")
    top_items = [
        "CREATOR -> PROCESSOR 1 (2 паралельні канали, delay = 1.0)",
        "PROCESSOR 1 розгалужується: 70% на PROCESSOR 2 та 30% на PROCESSOR 3",
        "PROCESSOR 2: 85% виходить у DISPOSE, а 15% повертається назад у PROCESSOR 1 (rework/зворотний зв'язок)",
        "PROCESSOR 3: 100% переходить у DISPOSE"
    ]
    for ti in top_items:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(ti)

    p = doc.add_paragraph()
    p.add_run("Результати верифікації завдання 6:")

    tbl6 = doc.add_table(rows=6, cols=4)
    tbl6.alignment = WD_TABLE_ALIGNMENT.CENTER
    h6 = ["Перехід", "Теоретична ймовірність", "Фактичні переходи (N / Всього)", "Емпірична ймовірність"]
    for j, h in enumerate(h6):
        c = tbl6.cell(0, j)
        set_cell_shading(c, "E2E8F0")
        set_cell_margins(c, 60, 60, 70, 70)
        c.paragraphs[0].add_run(h).bold = True
    
    d6 = [
        ["P1 -> P2 (Розгалуження)", "70.0%", "382 / 531", "71.9% (похибка 1.9%)"],
        ["P1 -> P3 (Розгалуження)", "30.0%", "149 / 531", "28.1% (похибка 1.9%)"],
        ["P2 -> DISPOSE", "85.0%", "333 / 373", "89.3% (похибка 4.3%)"],
        ["P2 -> P1 (Зворотний зв'язок)", "15.0%", "40 / 373", "10.7% (похибка 4.3%)"],
        ["P3 -> DISPOSE", "100.0%", "149 / 149", "100.0% (точний збіг)"]
    ]
    for i, row in enumerate(d6):
        for j, val in enumerate(row):
            c = tbl6.cell(i+1, j)
            set_cell_margins(c, 50, 50, 70, 70)
            c.paragraphs[0].add_run(val)

    p_bal = doc.add_paragraph()
    p_bal.paragraph_format.space_before = Pt(6)
    p_bal.add_run("Глобальне збереження кількості сутностей у мережі (Conservation Law):\n").bold = True
    p_bal.add_run("Створено CREATOR = 491 вимог.\n")
    p_bal.add_run("Вийшло через DISPOSE (482) + Загальна кількість відмов (8) + Залишилось у чергах і каналах (1) = 491 вимог.\n")
    p_bal.add_run("Підсумок: баланс сутностей виконується абсолютно точно, що доводить математичну коректність роботи розгалуження та зворотного зв'язку.")

    # --- Висновки ---
    add_h1("10. Висновки")
    concl = [
        "1. У ході виконання комп'ютерного практикуму успішно реалізовано об'єктно-орієнтований підхід до імітаційного моделювання дискретно-подійних систем мовою C#.",
        "2. Освоєно алгоритм просування модельного часу за принципом найближчої події (Next-Event Time Advance), що забезпечує високу обчислювальну швидкодію та точність моделювання без дискретизації за фіксованим кроком часу.",
        "3. Реалізовано обчислення середнього завантаження пристрою на основі інтегрування функцій стану каналів у часі.",
        "4. Побудовано та всебічно верифіковано трифазну послідовну систему масового обслуговування (Завдання 3 та 4). Визначено закономірності поведінки системи при зміні навантаження, виникненні 'вузьких місць' та обмеженнях місткості накопичувачів.",
        "5. Успішно розширено клас Process: додано повноцінну підтримку багатоканальності (Завдання 5) та гнучкої маршрутизації потоків за кумулятивними ймовірностями з підтримкою циклічних зворотних зв'язків (Завдання 6).",
        "6. Збережено 100% зворотну сумісність архітектури: базові завдання 1–3 функціонують без змін у коді на модифікованому ядрі класів.",
        "7. Математична адекватність моделей підтверджена точним збереженням балансу сутностей та збігом експериментальних результатів із теоретичними розрахунками теорії ймовірностей та теорії масового обслуговування."
    ]
    for c_text in concl:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(c_text)

    doc.save("Zvit_Praktikum_2.docx")
    print("Report generated successfully as Zvit_Praktikum_2.docx")

if __name__ == "__main__":
    create_report()
